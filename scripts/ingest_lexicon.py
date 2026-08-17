"""
Ingest the Volubilis Thai<->English(<->French) dictionary into the shared
`lexicon` table.

Source: the VOLUBILIS DATABASE edition (https://belisan-volubilis.blogspot.com/),
confirmed against v.26.2 (Jul. 2026), 114,577 entries. This is a different file
from the "Duo Max FRA" edition and has a different column order — do not reuse
column indices tuned for that file. Two header rows precede the data; columns
used here (0-indexed, verified directly against this file's actual header
row — cross-check any new source file's header before trusting these):
    1  EASYTHAI        -> lexicon.romanization
    4  THA (Thai)      -> lexicon.thai
    5  ENG (English)   -> lexicon.english (split on ';' into one row per sense)
    7  TYPE            -> lexicon.pos
    8  USAGE           -> lexicon.usage (register marker, e.g. "(obsol.)")
    9  SCIENT/abbrev.  -> lexicon.scientific_name (binomial, e.g. "Calotropis
                          gigantea R. Br." -> stored as "Calotropis gigantea")

This edition has no per-row Level (B/A1/A2/s) column, so lexicon.level is
always written as NULL here — do not fabricate a value. If a future source
file does carry a level column, wire it in as its own mapped column rather
than guessing from other fields.

A scientific binomial sometimes also appears as its own sense inside the ENG
cell (e.g. "Calotropis gigantea ; Crown flower" alongside
SCIENT/abbrev.="Calotropis gigantea R. Br."). Because it's already captured
in scientific_name, that sense is dropped from the split ENG list so it never
becomes its own gloss row — this is what stops "Calotropis gigantea" from
ever surfacing as a compound-breakdown gloss.

Idempotent: clears existing source="volubilis" rows before re-ingesting.
A real (non-dry-run) run refuses to proceed if it would replace more existing
rows than it has parsed as incoming replacements (see `--force`), and takes a
timestamped backup of the sqlite DB file before deleting anything.

Usage:
    python scripts/ingest_lexicon.py --source-file /path/to/VOLUBILIS.xlsx [--dry-run] [--force] [--batch-size N]

Flags:
    --source-file PATH  Path to the Volubilis .xlsx export (required).
    --dry-run           Parse and report counts without writing to DB.
    --force             Proceed even if incoming row count is less than existing
                         source="volubilis" row count.
    --batch-size N      Rows per bulk insert (default: 2000).
"""
import argparse
import asyncio
import re
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, "/app")

from loguru import logger
from sqlalchemy import delete, func, insert, select

from src.config.settings import get_settings
from src.db.database import get_session_factory
from src.db.models.lexicon import Lexicon

_HEADER_ROWS = 2
_THAI_COL = 4       # column 5, "THA (Thai)"
_ROMAN_COL = 1       # column 2, "EASYTHAI"
_ENGLISH_COL = 5       # column 6, "ENG (English)"
_POS_COL = 7       # column 8, "TYPE"
_USAGE_COL = 8       # column 9, "USAGE"
_SCIENT_COL = 9       # column 10, "SCIENT/abbrev."

# Leading "Genus species" from a SCIENT/abbrev. cell like "Calotropis gigantea
# R. Br." — drops the trailing taxonomic authorship abbreviation.
_SCIENTIFIC_NAME_RE = re.compile(r"^([A-Z][a-zA-Z\-]+\s+[a-z][a-zA-Z\-]+)")


def _clean_usage(cell) -> str | None:
    if not cell:
        return None
    text = str(cell).strip()
    if text.startswith("(") and text.endswith(")"):
        text = text[1:-1].strip()
    return text or None


def _parse_scientific_name(cell) -> str | None:
    if not cell:
        return None
    text = str(cell).strip()
    if not text:
        return None
    match = _SCIENTIFIC_NAME_RE.match(text)
    return match.group(1) if match else text


def _parse_rows(source_file: str):
    import openpyxl

    wb = openpyxl.load_workbook(source_file, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    for row in ws.iter_rows(min_row=_HEADER_ROWS + 1, values_only=True):
        thai = row[_THAI_COL]
        english_cell = row[_ENGLISH_COL]
        if not thai or not english_cell:
            continue

        thai = str(thai).strip()
        if not thai:
            continue

        roman = row[_ROMAN_COL]
        romanization = str(roman).strip() or None if roman else None
        pos_cell = row[_POS_COL]
        pos = str(pos_cell).strip() or None if pos_cell else None
        usage = _clean_usage(row[_USAGE_COL])
        scientific_name = _parse_scientific_name(row[_SCIENT_COL])

        senses = [
            sense.replace("\xa0", " ").strip()
            for sense in str(english_cell).split(";")
        ]
        senses = [sense for sense in senses if sense]
        if scientific_name:
            senses = [
                sense for sense in senses
                if sense.lower() != scientific_name.lower()
            ]

        for english in senses:
            yield {
                "thai": thai,
                "romanization": romanization,
                "english": english,
                "pos": pos,
                "usage": usage,
                "scientific_name": scientific_name,
                "level": None,
                "source": "volubilis",
            }


def _backup_sqlite_db() -> None:
    url = get_settings().database_url
    prefix = "sqlite+aiosqlite://"
    if not url.startswith(prefix):
        logger.info(f"Non-sqlite database_url ({url!r}) — skipping file backup")
        return

    db_path = Path(url[len(prefix):])
    if not db_path.exists():
        logger.warning(f"DB file {db_path} not found — skipping backup")
        return

    backup_dir = db_path.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    backup_path = backup_dir / f"{db_path.stem}.{stamp}{db_path.suffix}"
    shutil.copy2(db_path, backup_path)
    logger.info(f"Backed up {db_path} -> {backup_path}")


async def ingest(source_file: str, dry_run: bool, force: bool, batch_size: int) -> None:
    rows = list(_parse_rows(source_file))
    logger.info(f"Parsed {len(rows)} lexicon rows from {source_file}")

    if not rows:
        logger.info("Nothing to do.")
        return

    if dry_run:
        for sample in rows[:10]:
            logger.info(f"[dry-run] {sample}")
        logger.info(f"Dry run complete — {len(rows)} rows would be inserted")
        return

    factory = get_session_factory()
    async with factory() as db:
        existing = await db.scalar(
            select(func.count()).select_from(Lexicon).where(Lexicon.source == "volubilis")
        )
        if existing > 0 and existing > len(rows) and not force:
            logger.error(
                f"Refusing to overwrite {existing} existing source='volubilis' rows with only "
                f"{len(rows)} incoming rows — this looks like it would destroy real data. "
                f"Pass --force to override if this is intentional."
            )
            return

        _backup_sqlite_db()

        result = await db.execute(delete(Lexicon).where(Lexicon.source == "volubilis"))
        logger.info(f"Cleared {result.rowcount} existing source='volubilis' rows")

        for i in range(0, len(rows), batch_size):
            batch = rows[i : i + batch_size]
            await db.execute(insert(Lexicon), batch)
            logger.info(f"Inserted batch {i // batch_size + 1} ({min(i + batch_size, len(rows))}/{len(rows)})")

        await db.commit()

    logger.info(f"Done — ingested {len(rows)} lexicon rows")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-file", required=True, help="Path to the Volubilis .xlsx export")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--force", action="store_true", help="Override the existing-rows-vs-incoming-rows safety check")
    parser.add_argument("--batch-size", type=int, default=2000)
    args = parser.parse_args()

    asyncio.run(ingest(source_file=args.source_file, dry_run=args.dry_run, force=args.force, batch_size=args.batch_size))
